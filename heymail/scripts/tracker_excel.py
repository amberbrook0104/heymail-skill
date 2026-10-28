"""Portable template-based Excel exporter. No mail or credential access."""
from copy import copy
from datetime import datetime, timedelta
from pathlib import Path
import math


def export(rows, output, author):
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    template = Path(__file__).resolve().parents[1] / 'assets/templates/application-ledger.xlsx'
    book = load_workbook(template)
    sheet = book['投递台账']
    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            cell.value = None
    for index, record in enumerate(rows, 5):
        time = record.get('time')
        if isinstance(time, str):
            time = datetime.fromisoformat(time.replace('Z', '+00:00')).replace(tzinfo=None)
        progress = record.get('follow_up')
        if isinstance(progress, (int, float)):
            progress = (datetime(1899,12,30)+timedelta(days=progress)).strftime('%Y-%m-%d')
        elif isinstance(progress, dict):
            progress = '原有公式：'+progress.get('formula','')
        values = [record.get('company'),record.get('position'),record.get('location'),time,
                  record.get('agent'),record.get('recipients'),record.get('progress'),progress,
                  record.get('notes'),record.get('id')]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(index, col)
            cell._style = copy(sheet.cell(5, col)._style)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = 's'
            cell.fill = PatternFill('solid', fgColor=('FFFFFF' if index%2 else 'FBF8F5'))
        sheet.cell(index,4).number_format='yyyy-mm-dd hh:mm'
        sheet.cell(index,8).number_format='@'
        lines=1
        for col in (1,2,6,8,9):
            value=str(values[col-1] or '')
            width=sheet.column_dimensions[chr(64+col)].width
            count=sum(max(1,math.ceil(sum(2 if ord(c)>255 else 1 for c in line)/width)) for line in value.split('\n'))
            lines=max(lines,count)
        sheet.row_dimensions[index].height=max(22,lines*12+4)
    for name in list(sheet.tables):
        del sheet.tables[name]
    table=Table(displayName='ApplicationsTable',ref=f'A4:J{max(5,len(rows)+4)}')
    table.tableStyleInfo=TableStyleInfo(name='TableStyleLight1',showRowStripes=False)
    sheet.add_table(table)
    sheet.column_dimensions['J'].hidden=True
    sheet.freeze_panes='A5'
    book.properties.creator=author
    book.properties.lastModifiedBy=author
    book.save(output)
