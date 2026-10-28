"""Local ledger tests only. No SMTP or credential calls."""
from copy import deepcopy
from datetime import date
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('tracker',ROOT/'heymail/scripts/tracker.py')
tracker=importlib.util.module_from_spec(spec)
spec.loader.exec_module(tracker)
RUNTIME=bool(os.environ.get('HEYMAIL_TEST_NODE') and os.environ.get('HEYMAIL_TEST_MODULE_DIR'))
SETTINGS={'node':os.environ.get('HEYMAIL_TEST_NODE','node'),
          'module_dir':os.environ.get('HEYMAIL_TEST_MODULE_DIR','unused-test-modules'),
          'author':'示例作者 & 测试'}


class TrackerTests(unittest.TestCase):
    def setUp(self):
        base=ROOT/'work/tracker-test-temp'
        base.mkdir(parents=True,exist_ok=True)
        self.temp=tempfile.TemporaryDirectory(dir=base)
        Path(self.temp.name).resolve().relative_to(base.resolve())
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.settings={**SETTINGS,'root':str(self.root)}
        self.manifest={'to':'hr@example.com','subject':'测试岗位','body':'合成测试正文',
                       'agent':'Codex','application':{'id':'test-job','company':'示例公司','position':'研究实习生','location':'远程'}}
        self.receipt={'message_id':'<test-1@example.com>','digest':'synthetic','to':'hr@example.com',
                      'time':'2026-10-27T01:00:00+00:00','status':'accepted'}

    def test_agent_detection_and_beijing_time(self):
        self.assertEqual(tracker.agent_name(r'C:\Users\test\.workbuddy\skills\heymail\scripts\mail.py'),'WorkBuddy')
        self.assertEqual(tracker.agent_name(r'C:\Users\test\.codex\skills\heymail\scripts\mail.py'),'Codex')
        self.assertEqual(tracker.agent_name(r'C:\Users\test\Doubao\skills\mail.py'),'豆包')
        self.assertEqual(tracker.local_time('2026-10-27T16:30:00+00:00'),'2026-10-28T00:30:00')

    def test_same_message_deduplicates_and_agents_merge_per_job(self):
        tracker.record_event(self.settings,self.manifest,self.receipt,self.root/'receipt.json')
        tracker.record_event(self.settings,self.manifest,self.receipt,self.root/'receipt.json')
        other={**self.receipt,'to':'second@example.com','message_id':'<test-2@example.com>','time':'2026-10-28T01:00:00+00:00'}
        tracker.record_event(self.settings,self.manifest,other,self.root/'receipt2.json',agent='豆包')
        events=[tracker.read_json(p) for p in (self.root/'.events').glob('*.json')]
        self.assertEqual(len(events),2)
        rows=tracker.aggregate(events)
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['agent'],'Codex、豆包')
        self.assertEqual(rows[0]['time'],'2026-10-27T09:00:00')
        self.assertIn('2026-10-28',rows[0]['notes'])
        self.assertEqual(rows[0]['recipients'],'hr@example.com\nsecond@example.com')

    def test_actual_recipient_addresses_deduplicate_without_guessing(self):
        events=[{'recipient':'HR <hr@example.com>'},{'recipient':'HR@example.com'},
                {'recipient':None},{'recipient':''},{'recipient':'second@example.com'}]
        self.assertEqual(tracker.recipient_emails(events),'hr@example.com\nsecond@example.com')
        self.assertEqual(tracker.recipient_emails([{}]),'')

    def test_reader_maps_old_and_new_headers_to_same_manual_fields(self):
        legacy_headers=['下次跟进' if h=='最新进展' else h for h in tracker.HEADERS]
        for headers in [tracker.HEADERS, legacy_headers, [h for h in legacy_headers if h!='收件邮箱']]:
            values={'公司':'示例公司','岗位':'示例岗位','工作地点':'北京','投递时间':'2026-10-27',
                    '投递 Agent':'Codex','收件邮箱':'hr@example.com','当前进展':'面试中',
                    '下次跟进':'2026-11-02','最新进展':'已约11月2日面试','备注':'旧表中的人工备注','_record_id':'stable-old-job'}
            ns=tracker.NS['s'];sheet=ET.Element('{'+ns+'}worksheet');grid=ET.SubElement(sheet,'{'+ns+'}sheetData')
            for r,items in [(4,headers),(5,[values[h] for h in headers])]:
                row=ET.SubElement(grid,'{'+ns+'}row',{'r':str(r)})
                for i,value in enumerate(items):
                    cell=ET.SubElement(row,'{'+ns+'}c',{'r':chr(65+i)+str(r),'t':'inlineStr'})
                    ET.SubElement(ET.SubElement(cell,'{'+ns+'}is'),'{'+ns+'}t').text=value
            book=self.root/'legacy-or-new.xlsx'
            with zipfile.ZipFile(book,'w') as archive:archive.writestr('xl/worksheets/sheet1.xml',ET.tostring(sheet))
            cells=tracker.read_workbook_rows(book)['stable-old-job']
            self.assertEqual(cells['F'],'面试中')
            self.assertEqual(cells['G'],'已约11月2日面试' if '最新进展' in headers else '2026-11-02')
            self.assertEqual(cells['I'],'旧表中的人工备注')
            self.assertEqual(cells['K'],'hr@example.com' if '收件邮箱' in headers else None)

    def test_spreadsheet_failure_keeps_submission_event(self):
        with patch.object(tracker,'sync',side_effect=PermissionError()):
            result=tracker.record_submission({'tracker':self.settings,'display_name':'示例发件人'},self.manifest,self.receipt,self.root/'receipt.json',ROOT/'mail.py')
        self.assertFalse(result['ok'])
        self.assertTrue(result['recorded'])
        self.assertEqual(len(list((self.root/'.events').glob('*.json'))),1)

    def test_missing_metadata_does_not_invent_company(self):
        event=tracker.record_event(self.settings,{'to':'x@example.com','subject':'unknown'},self.receipt,self.root/'receipt.json')
        row=tracker.aggregate([event])[0]
        self.assertEqual(row['company'],'待确认')
        self.assertEqual(row['position'],'待确认')
        self.assertEqual(row['source']['label'],'邮件记录（JD待补）')

    @unittest.skipUnless(RUNTIME,'Set HEYMAIL_TEST_NODE and HEYMAIL_TEST_MODULE_DIR for Excel integration.')
    def test_refresh_preserves_manual_edits_after_row_sort(self):
        older={**self.manifest,'application':{**self.manifest['application'],'id':'older'}}
        tracker.record_event(self.settings,older,self.receipt,self.root/'old-receipt.json')
        newer={**self.manifest,'application':{**self.manifest['application'],'id':'newer','company':'第二家示例公司'}}
        tracker.record_event(self.settings,newer,{**self.receipt,'message_id':'<newer@example.com>','time':'2026-10-28T01:00:00+00:00'},self.root/'new-receipt.json',agent='豆包')
        tracker.sync(self.settings)
        book=self.root/'实习投递台账.xlsx'
        with zipfile.ZipFile(book) as z:data={n:z.read(n) for n in z.namelist()}
        ns=tracker.NS['s']; sheet=ET.fromstring(data['xl/worksheets/sheet1.xml'])
        grid=sheet.find('s:sheetData',tracker.NS)
        row5=next(r for r in grid if r.attrib['r']=='5')
        for col,value in [('A','手动确认的公司'),('G','面试中'),('H','已约11月2日面试，等待会议链接'),('I','手工备注，刷新后保留')]:
            cell=next((c for c in row5 if c.attrib['r']==col+'5'),None)
            if cell is None:cell=ET.SubElement(row5,'{'+ns+'}c',{'r':col+'5'})
            cell.clear();cell.set('r',col+'5')
            if isinstance(value,str):
                cell.set('t','inlineStr');ET.SubElement(ET.SubElement(cell,'{'+ns+'}is'),'{'+ns+'}t').text=value
            else:ET.SubElement(cell,'{'+ns+'}v').text=str(value)
        row6=next(r for r in grid if r.attrib['r']=='6')
        clones=[deepcopy(row6),deepcopy(row5)]
        indices=[list(grid).index(row5),list(grid).index(row6)]
        grid.remove(row5);grid.remove(row6)
        for newrow,r,index in zip(clones,[5,6],indices):
            newrow.set('r',str(r))
            for c in newrow:c.set('r',re.sub(r'\d+$',str(r),c.attrib['r']))
            grid.insert(index,newrow)
        for h in sheet.findall('s:hyperlinks/s:hyperlink',tracker.NS):
            if h.attrib['ref']=='H5':h.set('ref','H6')
            elif h.attrib['ref']=='H6':h.set('ref','H5')
        data['xl/worksheets/sheet1.xml']=ET.tostring(sheet,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(book,'w',zipfile.ZIP_DEFLATED) as z:
            for name,value in data.items():z.writestr(name,value)
        result=tracker.sync(self.settings)
        self.assertEqual(result['rows'],2)
        records=tracker.read_workbook_rows(book)
        self.assertEqual(list(records)[0],'newer')
        cells=records['newer']
        self.assertEqual(cells['A'],'手动确认的公司')
        self.assertEqual(cells['F'],'面试中')
        self.assertEqual(cells['I'],'手工备注，刷新后保留')
        self.assertEqual(cells['G'],'已约11月2日面试，等待会议链接')
        self.assertIsNone(cells['H'])
        self.assertEqual(cells['K'],'hr@example.com')
        with zipfile.ZipFile(book) as z:
            core=ET.fromstring(z.read('docProps/core.xml'))
            self.assertEqual(core.find('{http://purl.org/dc/elements/1.1/}creator').text,self.settings['author'])
            sheet=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
            self.assertEqual(sheet.find("s:cols/s:col[@min='10']",tracker.NS).attrib['hidden'],'1')
            self.assertEqual(len(list(sheet.find("s:sheetData/s:row[@r='4']",tracker.NS))),10)
            self.assertFalse(sheet.findall('s:hyperlinks/s:hyperlink',tracker.NS))
            self.assertFalse(any('TODAY()' in ''.join(rule.itertext()) for rule in sheet.findall('s:conditionalFormatting',tracker.NS)))

    def test_legacy_dates_remain_readable_in_latest_progress(self):
        serial=float((date(2026,11,2)-date(1899,12,30)).days)
        self.assertEqual(tracker.latest_progress_text(serial,legacy_date=True),'2026-11-02')
        self.assertEqual(tracker.latest_progress_text('HR已回复，等待面试通知'),'HR已回复，等待面试通知')

    def test_numeric_and_iso_dates_sort_consistently(self):
        values=['2026-10-27T00:00:00',float((date(2026,10,28)-date(1899,12,30)).days),'2026-10-26T00:00:00Z']
        ordered=sorted(values,key=tracker.time_sort_key,reverse=True)
        self.assertIsInstance(ordered[0],float)
        self.assertEqual(ordered[1],values[0])

    def test_python_backend_preserves_edits_and_native_features(self):
        from openpyxl import load_workbook
        settings={'root':str(self.root),'backend':'python','author':'Synthetic Test'}
        tracker.record_event(settings,self.manifest,self.receipt,self.root/'receipt.json',agent='WorkBuddy')
        tracker.sync(settings)
        output=self.root/'实习投递台账.xlsx'
        book=load_workbook(output);sheet=book['投递台账']
        self.assertEqual(sheet['E5'].value,'WorkBuddy')
        self.assertEqual(sheet['F5'].value,'hr@example.com')
        self.assertEqual(sheet.freeze_panes,'A5')
        self.assertTrue(sheet.column_dimensions['J'].hidden)
        self.assertEqual(len(sheet._images),2)
        sheet['G5']='面试中';sheet['H5']='HR已回复，等待面试';sheet['I5']='自填备注'
        book.save(output)
        tracker.sync(settings)
        cells=tracker.read_workbook_rows(output)['test-job']
        self.assertEqual(cells['F'],'面试中')
        self.assertEqual(cells['G'],'HR已回复，等待面试')
        self.assertEqual(cells['I'],'自填备注')
        book=load_workbook(output);sheet=book['投递台账']
        self.assertEqual(sheet['H5'].number_format,'@')
        self.assertEqual(sheet.tables['ApplicationsTable'].ref,'A4:J5')
        self.assertEqual(book.properties.creator,'Synthetic Test')

    @unittest.skipUnless(os.name=='nt','Windows shortcut feature.')
    def test_shortcut_reuses_same_target_and_preserves_conflicting_shortcut(self):
        desktop=self.root/'desktop';desktop.mkdir()
        book=self.root/'实习投递台账.xlsx';book.write_bytes(b'synthetic fixture; never opened')
        first=tracker.create_desktop_shortcut(self.settings,desktop)
        second=tracker.create_desktop_shortcut(self.settings,desktop)
        self.assertEqual(first['shortcut'],second['shortcut'])
        self.assertEqual(Path(first['target']),book)
        other=self.root/'other';other.mkdir();(other/'实习投递台账.xlsx').write_bytes(b'second fixture')
        third=tracker.create_desktop_shortcut({**self.settings,'root':str(other)},desktop)
        self.assertNotEqual(first['shortcut'],third['shortcut'])
        self.assertTrue(Path(first['shortcut']).is_file())

    @unittest.skipUnless(os.name=='nt','Windows installer feature.')
    def test_first_install_and_replace_preserve_external_profile(self):
        destination=self.root/'installed/heymail';backup=self.root/'backups'
        config=self.root/'local-account-config.json';config.write_text('{"keep":"synthetic-profile"}',encoding='utf-8')
        command=['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(ROOT/'install.ps1'),
                 '-Destination',str(destination),'-BackupDirectory',str(backup)]
        installed=subprocess.run(command,capture_output=True)
        self.assertEqual(installed.returncode,0,installed.stderr.decode('utf-8',errors='replace'))
        (destination/'SKILL.md').write_text('synthetic old version',encoding='utf-8')
        updated=subprocess.run(command+['-Replace'],capture_output=True)
        self.assertEqual(updated.returncode,0,updated.stderr.decode('utf-8',errors='replace'))
        self.assertEqual((destination/'SKILL.md').read_bytes(),(ROOT/'heymail/SKILL.md').read_bytes())
        self.assertTrue(any(p.read_text(encoding='utf-8')=='synthetic old version' for p in backup.rglob('SKILL.md')))
        self.assertEqual(config.read_text(encoding='utf-8'),'{"keep":"synthetic-profile"}')
        self.assertFalse(list(destination.rglob('*.pyc')))


if __name__=='__main__':unittest.main()
